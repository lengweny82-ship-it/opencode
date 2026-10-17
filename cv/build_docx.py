# -*- coding: utf-8 -*-
"""Build the Arabic and English CVs as editable Word (.docx) files."""

import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from content import CONTENT

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "CV")
PHOTO = os.path.join(HERE, "assets", "photo_circle_docx.png")

NAVY = "102440"
GOLD = "A28040"
GOLD_LIGHT = "C6AA70"
INK = "262C34"
MUTED = "6E7682"
RULE = "D6D8DC"
CHIP_BG = "F6F8FB"
CHIP_BORDER = "DFE3E9"
BAND_TEXT = "CED5DF"
WHITE = "FFFFFF"

AR_FONT = "Arial"
EN_FONT = "Calibri"
EN_DISPLAY = "Georgia"

TBLPR_SEQ = ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
             "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
             "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout",
             "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"]

TCPR_SEQ = ["w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders",
            "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
            "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge",
            "w:tcPrChange"]

PPR_SEQ = ["w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
           "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
           "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
           "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
           "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
           "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
           "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
           "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange"]

RPR_SEQ = ["w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps",
           "w:smallCaps", "w:strike", "w:dstrike", "w:outline", "w:shadow", "w:emboss",
           "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden",
           "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs",
           "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText",
           "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
           "w:specVanish", "w:oMath"]


def ordered_insert(parent, el, seq):
    """Insert `el` into `parent` at its schema-correct position."""
    tag = el.tag
    name = None
    for candidate in seq:
        if qn(candidate) == tag:
            name = candidate
            break
    if name is None:
        parent.append(el)
        return el
    successors = seq[seq.index(name) + 1:]
    parent.insert_element_before(el, *successors)
    return el


def child(parent, tag, seq):
    """Return existing direct child by tag, removing duplicates."""
    found = parent.findall(qn(tag))
    for extra in found[1:]:
        parent.remove(extra)
    if found:
        return found[0]
    return None


def get_or_make(parent, tag, seq):
    existing = child(parent, tag, seq)
    if existing is not None:
        return existing
    return ordered_insert(parent, OxmlElement(tag), seq)


# ---------------------------------------------------------------- formatting
def set_bidi(p):
    pPr = p._p.get_or_add_pPr()
    get_or_make(pPr, "w:bidi", PPR_SEQ)


def shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = get_or_make(pPr, "w:shd", PPR_SEQ)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def para_border(p, side="bottom", color=GOLD, sz=8, space=2):
    pPr = p._p.get_or_add_pPr()
    pBdr = get_or_make(pPr, "w:pBdr", PPR_SEQ)
    old = pBdr.find(qn(f"w:{side}"))
    if old is not None:
        pBdr.remove(old)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pBdr.append(el)


def style_run(run, size, color, bold=False, font=AR_FONT, cs_font=None,
              rtl=False, spacing=None, caps=False):
    run.font.size = Pt(size)
    if bold:
        run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = font
    rPr = run._element.get_or_add_rPr()

    rFonts = rPr.find(qn("w:rFonts"))
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:cs"), cs_font or font)

    szCs = get_or_make(rPr, "w:szCs", RPR_SEQ)
    szCs.set(qn("w:val"), str(int(round(size * 2))))

    if bold:
        bCs = get_or_make(rPr, "w:bCs", RPR_SEQ)
        bCs.set(qn("w:val"), "1")

    if rtl:
        get_or_make(rPr, "w:rtl", RPR_SEQ)

    if spacing is not None:
        el = get_or_make(rPr, "w:spacing", RPR_SEQ)
        el.set(qn("w:val"), str(int(round(spacing * 20))))
    if caps:
        get_or_make(rPr, "w:caps", RPR_SEQ)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = get_or_make(tcPr, "w:shd", TCPR_SEQ)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:tcMar"))
    if old is not None:
        tcPr.remove(old)
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    ordered_insert(tcPr, mar, TCPR_SEQ)


def set_cell_borders(cell, color, sz=4, sides=("top", "left", "bottom", "right")):
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:tcBorders"))
    if old is not None:
        tcPr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for side in sides:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    ordered_insert(tcPr, borders, TCPR_SEQ)


def set_cell_valign(cell, val="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    el = get_or_make(tcPr, "w:vAlign", TCPR_SEQ)
    el.set(qn("w:val"), val)


def table_no_borders(table):
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    ordered_insert(tblPr, borders, TBLPR_SEQ)


def table_fixed(table, widths_cm):
    table.autofit = False  # python-docx writes w:tblLayout fixed
    tblPr = table._tbl.tblPr
    total = int(round(sum(widths_cm) * 567))
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        ordered_insert(tblPr, tblW, TBLPR_SEQ)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(total))

    for row in table.rows:
        for idx, w in enumerate(widths_cm):
            row.cells[idx].width = Cm(w)
    grid = table._tbl.tblGrid
    for idx, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(round(widths_cm[idx] * 567))))


def table_bidi(table):
    tblPr = table._tbl.tblPr
    get_or_make(tblPr, "w:bidiVisual", TBLPR_SEQ)


def table_center(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER


def add_run(p, text, size, color, bold=False, font=AR_FONT, cs_font=None,
            rtl=False, spacing=None, caps=False):
    r = p.add_run(text)
    style_run(r, size, color, bold=bold, font=font, cs_font=cs_font,
              rtl=rtl, spacing=spacing, caps=caps)
    return r


def space(p, before=0, after=0, line=1.12):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def set_indent(p, start_twips, hanging_twips, rtl):
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn("w:ind"))
    if old is not None:
        pPr.remove(old)
    ind = OxmlElement("w:ind")
    side = "w:right" if rtl else "w:left"
    ind.set(qn(side), str(start_twips))
    ind.set(qn("w:hanging"), str(hanging_twips))
    ordered_insert(pPr, ind, PPR_SEQ)


def keep_next(p):
    p.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------- document
def build(lang, path):
    c = CONTENT[lang]
    ar = lang == "ar"
    body_font = AR_FONT if ar else EN_FONT
    align_start = WD_ALIGN_PARAGRAPH.RIGHT if ar else WD_ALIGN_PARAGRAPH.LEFT

    photo_path = PHOTO
    if not os.path.exists(photo_path):
        raise SystemExit(f"missing photo: {photo_path}")

    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = Cm(0.8)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.9)
    sec.right_margin = Cm(1.9)
    sec.header_distance = Cm(0.5)
    sec.footer_distance = Cm(0.8)
    text_w = 21.0 - 3.8  # 17.2 cm

    normal = doc.styles["Normal"]
    normal.font.name = body_font
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    rfonts.set(qn("w:cs"), body_font)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    # ---------------- header band ----------------
    photo_w = 2.3          # cm (circular portrait)
    photo_cell_w = photo_w + 0.9
    text_cell_w = text_w - photo_cell_w

    band = doc.add_table(rows=1, cols=2)
    table_no_borders(band)
    table_fixed(band, [text_cell_w, photo_cell_w])
    table_center(band)
    if ar:
        table_bidi(band)

    # logical order: text cell first, photo second -> mirrors correctly in RTL
    text_cell = band.cell(0, 0)
    set_cell_shading(text_cell, NAVY)
    set_cell_margins(text_cell, top=250, bottom=210, left=300, right=300)
    set_cell_borders(text_cell, GOLD, sz=18, sides=("bottom",))
    set_cell_valign(text_cell, "center")

    p = text_cell.paragraphs[0]
    if ar:
        set_bidi(p)
    p.alignment = align_start
    space(p, before=0, after=2, line=1.0)
    if ar:
        add_run(p, c["name"], 22, WHITE, bold=True, font=AR_FONT, cs_font=AR_FONT, rtl=True)
    else:
        add_run(p, c["name"], 22, WHITE, bold=True, font=EN_DISPLAY, cs_font=EN_DISPLAY)

    p = text_cell.add_paragraph()
    if ar:
        set_bidi(p)
    p.alignment = align_start
    space(p, before=0, after=5, line=1.0)
    if ar:
        add_run(p, c["title"], 11.5, GOLD_LIGHT, bold=True, font=AR_FONT,
                cs_font=AR_FONT, rtl=True)
    else:
        add_run(p, c["title"], 9, GOLD_LIGHT, bold=True, font=EN_FONT, spacing=1.4, caps=True)

    p = text_cell.add_paragraph()
    if ar:
        set_bidi(p)
    p.alignment = align_start
    space(p, before=0, after=0, line=1.0)
    add_run(p, f'{c["phone_label"]}: {c["phone"]}', 9.5, BAND_TEXT,
            font=body_font, cs_font=AR_FONT, rtl=ar)

    # photo cell
    photo_cell = band.cell(0, 1)
    set_cell_shading(photo_cell, NAVY)
    set_cell_margins(photo_cell, top=150, bottom=150, left=140, right=140)
    set_cell_borders(photo_cell, GOLD, sz=18, sides=("bottom",))
    set_cell_valign(photo_cell, "center")
    p = photo_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(p, before=0, after=0, line=1.0)
    run = p.add_run()
    run.add_picture(photo_path, width=Cm(photo_w))

    spacer = doc.add_paragraph()
    space(spacer, before=0, after=0, line=1.0)
    spacer.paragraph_format.line_spacing = Pt(1)

    # ---------------- section helper ----------------
    def section(title):
        p = doc.add_paragraph()
        if ar:
            set_bidi(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            space(p, before=13, after=4, line=1.0)
            add_run(p, "|", 15, GOLD, bold=True, font=AR_FONT, cs_font=AR_FONT, rtl=True)
            add_run(p, "  " + title, 13, NAVY, bold=True, font=AR_FONT,
                    cs_font=AR_FONT, rtl=True)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            space(p, before=13, after=4, line=1.0)
            add_run(p, "|", 13.5, GOLD, bold=True, font=EN_DISPLAY)
            add_run(p, "  " + title, 11, NAVY, bold=True, font=EN_DISPLAY,
                    spacing=1.1, caps=True)
        para_border(p, "bottom", GOLD, sz=8, space=4)
        keep_next(p)
        return p

    def body_para(text, after=4):
        p = doc.add_paragraph()
        if ar:
            set_bidi(p)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        space(p, before=0, after=after, line=1.10)
        add_run(p, text, 10.2, INK, font=body_font, cs_font=AR_FONT, rtl=ar)
        return p

    def bullet(text, after=3, size=10.2):
        p = doc.add_paragraph()
        if ar:
            set_bidi(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        space(p, before=0, after=after, line=1.10)
        set_indent(p, 340, 340, ar)
        add_run(p, "▪  ", size, GOLD, bold=True, font=AR_FONT, cs_font=AR_FONT, rtl=ar)
        add_run(p, text, size, INK, font=body_font, cs_font=AR_FONT, rtl=ar)
        return p

    # ---------------- summary ----------------
    section(c["sections"]["summary"])
    body_para(c["summary"])

    # ---------------- practice areas ----------------
    section(c["sections"]["practice"])
    items = c["practice"]
    half = (len(items) + 1) // 2
    cols = [items[:half], items[half:]]
    t = doc.add_table(rows=half, cols=2)
    table_no_borders(t)
    table_fixed(t, [text_w / 2, text_w / 2])
    if ar:
        table_bidi(t)
    for r in range(half):
        for col in range(2):
            cell = t.cell(r, col)
            set_cell_margins(cell, top=25, bottom=25, left=0, right=0)
            p = cell.paragraphs[0]
            if ar:
                set_bidi(p)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            space(p, before=0, after=0, line=1.1)
            add_run(p, "▪  ", 10.2, GOLD, bold=True, font=AR_FONT, cs_font=AR_FONT, rtl=ar)
            add_run(p, cols[col][r], 10.2, INK, font=body_font, cs_font=AR_FONT, rtl=ar)

    # ---------------- experience ----------------
    section(c["sections"]["experience"])
    t = doc.add_table(rows=1, cols=2)
    table_no_borders(t)
    table_fixed(t, [text_w * 0.62, text_w * 0.38])
    if ar:
        table_bidi(t)
    cell = t.cell(0, 0)
    set_cell_margins(cell, top=0, bottom=0, left=0, right=0)
    p = cell.paragraphs[0]
    if ar:
        set_bidi(p)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    space(p, before=0, after=0, line=1.0)
    add_run(p, c["role"], 11.5, NAVY, bold=True, font=body_font, cs_font=AR_FONT, rtl=ar)
    cell = t.cell(0, 1)
    set_cell_margins(cell, top=0, bottom=0, left=0, right=0)
    p = cell.paragraphs[0]
    if ar:
        set_bidi(p)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    space(p, before=0, after=0, line=1.0)
    add_run(p, c["period"], 10, GOLD, bold=True, font=body_font, cs_font=AR_FONT, rtl=ar)

    for b in c["experience"]:
        bullet(b)

    # ---------------- education ----------------
    section(c["sections"]["education"])
    bullet(c["education_degree"], after=0)

    # ---------------- skills ----------------
    section(c["sections"]["skills"])
    skills = c["skills"]
    per_row = 3
    rows = (len(skills) + per_row - 1) // per_row
    gap = 0.35
    chip_w = (text_w - gap * (per_row - 1)) / per_row
    t = doc.add_table(rows=rows, cols=per_row)
    table_no_borders(t)
    table_fixed(t, [chip_w] * per_row)
    table_center(t)
    if ar:
        table_bidi(t)
    for r in range(rows):
        for col in range(per_row):
            idx = r * per_row + col
            cell = t.cell(r, col)
            set_cell_borders(cell, CHIP_BORDER, sz=4)
            set_cell_shading(cell, CHIP_BG)
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            set_cell_valign(cell, "center")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            space(p, before=0, after=0, line=1.0)
            if idx < len(skills):
                add_run(p, skills[idx], 9.0, "1E3452", font=body_font,
                        cs_font=AR_FONT, rtl=ar)

    # ---------------- courses ----------------
    section(c["sections"]["courses"])
    for i, course in enumerate(c["courses"]):
        bullet(course, after=3 if i < len(c["courses"]) - 1 else 0)

    # ---------------- footer ----------------
    footer = sec.footer
    fp = footer.paragraphs[0]
    if ar:
        set_bidi(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(fp, before=0, after=0, line=1.0)
    para_border(fp, "top", RULE, sz=4, space=6)
    add_run(fp, f'{c["phone_label"]}  {c["phone"]}   ·   {c["name"]}', 8, MUTED,
            font=body_font, cs_font=AR_FONT, rtl=ar)

    doc.core_properties.title = f"{c['name']} — CV"
    doc.core_properties.author = c["name"]
    doc.core_properties.subject = "Curriculum Vitae"
    doc.save(path)
    print("wrote", path)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    build("ar", os.path.join(OUT, "Ahmed_AbdelNaby_CV_AR.docx"))
    build("en", os.path.join(OUT, "Ahmed_AbdelNaby_CV_EN.docx"))
